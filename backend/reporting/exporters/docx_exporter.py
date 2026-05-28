import os
try:
    from docx import Document
    from docx.shared import Pt, Inches
    from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
except ImportError:
    Document = None

class DOCXExporter:
    def __init__(self):
        pass

    def export(self, report_data: dict, output_path: str) -> bool:
        if Document is None:
            raise ImportError("python-docx is not installed.")

        doc = Document()
        
        # Title page
        title = doc.add_heading(report_data.get('title', 'Vulnerability Scan Report'), 0)
        title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        
        doc.add_paragraph(f"Generated: {report_data.get('generated_timestamp', 'Unknown')}").alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        doc.add_page_break()

        # Executive Summary
        if 'executive_summary' in report_data:
            doc.add_heading('Executive Summary', level=1)
            doc.add_paragraph(report_data['executive_summary'])
            doc.add_page_break()

        # Detailed Findings
        if 'detailed_findings' in report_data and report_data['detailed_findings']:
            doc.add_heading('Detailed Findings', level=1)
            for finding_raw in report_data['detailed_findings']:
                if not isinstance(finding_raw, dict):
                    finding = getattr(finding_raw, '__dict__', {})
                else:
                    finding = finding_raw
                doc.add_heading(f"[{finding.get('severity', 'UNKNOWN').upper()}] {finding.get('name', 'Unnamed')}", level=2)
                
                table = doc.add_table(rows=4, cols=2)
                table.style = 'Table Grid'
                
                table.cell(0, 0).text = "Host"
                table.cell(0, 1).text = finding.get('host', 'N/A')
                
                table.cell(1, 0).text = "Port/Service"
                table.cell(1, 1).text = f"{finding.get('port', 'N/A')} / {finding.get('service', 'N/A')}"
                
                table.cell(2, 0).text = "CVSS Score"
                table.cell(2, 1).text = str(finding.get('cvss', 'N/A'))
                
                table.cell(3, 0).text = "CVE"
                table.cell(3, 1).text = finding.get('cve', 'N/A')
                
                doc.add_heading('Description', level=3)
                doc.add_paragraph(finding.get('description', ''))
                
                if finding.get('ai_explanation'):
                    doc.add_heading('AI Explanation', level=3)
                    doc.add_paragraph(finding.get('ai_explanation', ''))
                    
                doc.add_heading('Remediation', level=3)
                doc.add_paragraph(finding.get('remediation', ''))
                
                doc.add_heading('Evidence', level=3)
                doc.add_paragraph(finding.get('evidence', ''))
                
                doc.add_paragraph("_" * 50) # separator
                
        doc.save(output_path)
        return True
